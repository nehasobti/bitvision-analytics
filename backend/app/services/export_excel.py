from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.i18n import dimension_label, field_labels, filter_label, t
from app.services.mining_report import sections, settings_rows
from app.services.my_tenders_report import columns as my_columns, filter_rows, summary_rows
from app.services.olap import STATS

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
TOTAL_FONT = Font(bold=True)

MONEY = '#,##0.00 "€"'
PERCENT = '0.000"%"'
DATE = "dd/mm/yyyy"


def _write_table(ws, headers: list[str], rows: list[list], formats: list[str | None]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in rows:
        ws.append(row)

    for col_idx, fmt in enumerate(formats, start=1):
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                cell.number_format = fmt

    for col_idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[col_idx - 1] or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 60)

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions


def _measure_format(measure: str) -> str:
    return MONEY if measure == "amount_awarded" else PERCENT


def build_excel(data: list[dict], olap: dict, filters: dict) -> bytes:
    wb = Workbook()

    # Sheet 1: filtered data
    ws = wb.active
    ws.title = t("report.sheet_data")
    labels = field_labels()
    columns = list(labels)
    formats = {
        "amount_awarded": MONEY,
        "downside_amount": PERCENT,
        "publication_date": DATE,
        "award_date": DATE,
    }
    _write_table(
        ws,
        [labels[c] for c in columns],
        [[row.get(c) for c in columns] for row in data],
        [formats.get(c) for c in columns],
    )

    # Sheet 2: OLAP analysis
    ws = wb.create_sheet(t("report.sheet_analysis"))
    group_by, measures = olap["group_by"], olap["measures"]
    stat_keys = [f"{m}_{s}" for m in measures for s in STATS]
    headers = (
        [dimension_label(g) for g in group_by]
        + [t("stats.count")]
        + [f"{t(f'fields.{m}')} – {t(f'stats.{s}')}" for m in measures for s in STATS]
    )
    rows = [[r.get(g) for g in group_by] + [r["count"]] + [r.get(k) for k in stat_keys] for r in olap["rows"]]
    totals = olap["totals"]
    rows.append([t("olap.total").upper()] * max(len(group_by), 1) + [totals["count"]] + [totals.get(k) for k in stat_keys])
    if not group_by:
        headers = [""] + headers
    fmts = [None] * max(len(group_by), 1) + ["0"] + [_measure_format(m) for m in measures for _ in STATS]
    _write_table(ws, headers, rows, fmts)
    for cell in ws[ws.max_row]:
        cell.font = TOTAL_FONT

    # Sheet 3: what was exported
    ws = wb.create_sheet(t("report.sheet_filters"))
    ws.append([t("report.generated"), datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([t("report.rows_exported"), len(data)])
    ws.append([t("report.grouped_by"), ", ".join(dimension_label(g) for g in group_by) or "—"])
    for key, value in filters.items():
        ws.append([filter_label(key), value.strftime("%d/%m/%Y") if isinstance(value, date) else value])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_KIND_FORMATS = {"money": MONEY, "percent": PERCENT, "date": DATE, "int": "0", "number": "0.00", "points": "+0.000;-0.000;0.000"}


def _cell(value, kind: str):
    if kind == "list":
        return ", ".join(value or [])
    if kind == "direction":
        return t(f"mining.{value}") if value else ""
    if kind == "status":
        return t(f"status.{value}") if value else ""
    if kind == "result":
        return t(f"my.{value}") if value else ""
    return value


def build_mining_excel(result: dict, filters: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    for section in sections(result):
        # Excel sheet names: max 31 characters.
        ws = wb.create_sheet(t(f"mining.sheet_{section['key']}")[:31])
        cols = section["columns"]
        _write_table(
            ws,
            [label for _, label, _ in cols],
            [[_cell(r.get(key), kind) for key, _, kind in cols] for r in section["rows"]],
            [_KIND_FORMATS.get(kind) for _, _, kind in cols],
        )
        if not section["rows"]:
            ws.append([t("mining.none")])
            ws.cell(ws.max_row, 1).font = Font(italic=True)
        if section["hint"]:
            ws.insert_rows(1)
            ws["A1"] = section["hint"]
            ws["A1"].font = Font(italic=True)
            ws.freeze_panes = "A3"
            # insert_rows does not move the filter range: set it again on the table (header now on row 2).
            if section["rows"]:
                ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{ws.max_row}"

    ws = wb.create_sheet(t("report.sheet_filters"))
    ws.append([t("report.generated"), datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([t("mining.records", count=result["records"]), ""])
    for label, value in settings_rows(result):
        ws.append([label, value])
    for key, value in filters.items():
        ws.append([filter_label(key), value.strftime("%d/%m/%Y") if isinstance(value, date) else value])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 40

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_my_tenders_excel(result: dict, filters: dict, azienda_name: str | None) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = t("my.sheet")[:31]
    cols = my_columns()
    _write_table(
        ws,
        [label for _, label, _ in cols],
        [[_cell(r.get(key), kind) for key, _, kind in cols] for r in result["rows"]],
        [_KIND_FORMATS.get(kind) for _, _, kind in cols],
    )
    if not result["rows"]:
        ws.append([t("my.none")])
        ws.cell(ws.max_row, 1).font = Font(italic=True)

    ws = wb.create_sheet(t("report.sheet_filters"))
    ws.append([t("report.generated"), datetime.now().strftime("%d/%m/%Y %H:%M")])
    for label, value in filter_rows(filters, azienda_name):
        ws.append([label, value])
    ws.append([])
    for label, value, kind in summary_rows(result["summary"]):
        ws.append([label, value])
        ws.cell(ws.max_row, 2).number_format = _KIND_FORMATS.get(kind, "General")
    ws.append([])
    ws.append([t("my.hint")])
    ws.cell(ws.max_row, 1).font = Font(italic=True)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 40

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
